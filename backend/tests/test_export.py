import csv
import io
from datetime import date

from app.models.azienda import Azienda
from app.models.lavoratore import Lavoratore
from app.services.export import IMPORTAZIONE_HEADER, genera_csv_importazione, genera_xlsx_allegato1


def _azienda():
    return Azienda(
        id=1,
        ragione_sociale="ACME SRL",
        partita_iva="12345678901",
        ateco="10.11",
        email_aziendale="acme@example.com",
        regione="LOMBARDIA",
    )


def _lavoratore(mansione="Operaio edile"):
    return Lavoratore(
        id=1,
        azienda_id=1,
        nome="Mario",
        cognome="Rossi",
        email="acme@example.com",
        data_nascita=date(1990, 5, 20),
        luogo_nascita="Milano",
        codice_fiscale="RSSMRA90E20F205X",
        mansione=mansione,
        area="LOMBARDIA",
    )


def test_header_esatto_come_template():
    csv_content = genera_csv_importazione(_azienda(), [])
    prima_riga = csv_content.splitlines()[0]
    assert prima_riga.split(",") == IMPORTAZIONE_HEADER


def test_una_riga_per_ogni_corso_assegnato():
    azienda = _azienda()
    lavoratore = _lavoratore("Operaio edile")  # -> 2 corsi: gen_lav, spec_alto

    csv_content = genera_csv_importazione(azienda, [lavoratore])
    righe = list(csv.reader(io.StringIO(csv_content)))

    assert len(righe) == 3  # header + 2 corsi
    for riga in righe[1:]:
        assert riga[0] == "Mario"
        assert riga[1] == "Rossi"
        assert riga[3] == "20/05/1990"
        assert riga[9] == "12345678901"
        assert riga[10].startswith("MANCA_SKU:")


def test_lavoratore_senza_corso_matchato_resta_in_export():
    azienda = _azienda()
    lavoratore = _lavoratore("Astronauta")

    csv_content = genera_csv_importazione(azienda, [lavoratore])
    righe = list(csv.reader(io.StringIO(csv_content)))

    assert len(righe) == 2  # header + 1 riga senza corso
    assert righe[1][10] == ""


def test_xlsx_allegato1_contiene_intestazione_e_lavoratori():
    from openpyxl import load_workbook

    azienda = _azienda()
    lavoratore = _lavoratore("Commesso")

    contenuto = genera_xlsx_allegato1(azienda, [lavoratore], datore_di_lavoro="Luigi Bianchi")
    wb = load_workbook(io.BytesIO(contenuto))
    ws = wb.active

    assert [c.value for c in ws[1]] == ["LAVORATORI", "MANSIONE", "AZIENDA", "ATECO", "DATORE DI LAVORO"]
    assert [c.value for c in ws[2]] == ["Rossi Mario", "Commesso", "ACME SRL", "10.11", "Luigi Bianchi"]
